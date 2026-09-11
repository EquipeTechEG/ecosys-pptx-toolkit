"""Utilitaires de remplissage du template Ecosys.

A importer dans le script de generation plutot que de les reecrire :

    import sys; sys.path.insert(0, 'chemin/vers/skill/scripts')
    from aide_pptx import forme, ecrire, ph, remplir_sommaire
"""
import copy

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches


def walk(shapes):
    """Parcourt toutes les formes, y compris celles imbriquees dans des Groupes."""
    for sh in shapes:
        yield sh
        if sh.__class__.__name__ == 'GroupShape':
            yield from walk(sh.shapes)


def forme(slide, nom):
    """Retrouve une forme par son nom (champ_... / visuel_...), Groupes inclus."""
    for sh in walk(slide.shapes):
        if sh.name == nom:
            return sh
    dispo = sorted(sh.name for sh in walk(slide.shapes) if sh.name.startswith('champ_'))
    raise KeyError(f'forme {nom!r} introuvable. Champs disponibles : {dispo}')


def ecrire(shape, lignes, style_de=0):
    """Ecrit une ligne par paragraphe en preservant la mise en forme du template."""
    if isinstance(lignes, str):
        lignes = [lignes]
    tf = shape.text_frame
    ref = tf.paragraphs[min(style_de, len(tf.paragraphs) - 1)]._p

    while len(tf.paragraphs) < len(lignes):
        tf._txBody.append(copy.deepcopy(ref))
    for p, texte in zip(tf.paragraphs, lignes):
        if p.runs:
            p.runs[0].text = texte
            for r in p.runs[1:]:
                r._r.getparent().remove(r._r)
        else:
            p.add_run().text = texte
    for p in list(tf.paragraphs)[len(lignes):]:
        p._p.getparent().remove(p._p)


def ph(slide, idx, lignes):
    """Ecrit dans un placeholder natif, cible par son idx."""
    ecrire(slide.placeholders[idx], lignes)


SOMMAIRE_LIBELLES = (0, 12, 13, 14, 15)
SOMMAIRE_NUMEROS = (10, 16, 17, 18, 19)
TRAITS_Y = (1.00, 2.13, 3.27, 4.40, 5.58)


def remplir_sommaire(slide, chapitres):
    """Remplit sommaire_liste avec N chapitres (1 a 5) et masque les traits en trop."""
    n = len(chapitres)
    if not 1 <= n <= 5:
        raise ValueError(f'{n} chapitres : sommaire_liste en accepte 1 a 5, '
                         'utiliser sommaire_grille (6 sections) ou regrouper')

    for i, libelle in enumerate(chapitres):
        ph(slide, SOMMAIRE_LIBELLES[i], libelle)
        ph(slide, SOMMAIRE_NUMEROS[i], str(i + 1))

    for i in range(n, 5):
        for idx in (SOMMAIRE_LIBELLES[i], SOMMAIRE_NUMEROS[i]):
            slide.placeholders[idx].text_frame.text = ''

    if n < 5:
        haut = TRAITS_Y[n]
        masque = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(5.90), Inches(haut - 0.06),
            Inches(0.35), Inches(TRAITS_Y[4] + 0.84 + 0.06 - (haut - 0.06)))
        masque.fill.solid()
        masque.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        masque.line.fill.background()
        masque.shadow.inherit = False
        masque.name = 'masque_traits_sommaire'
    return n


import copy as _copy

VERT = (0x10, 0xBC, 0x4B)
BLEU = (0x00, 0xA3, 0xE9)
TITRES_NOMMES = {'champ_titre', 'champ_titre_annexe'}

_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
_NS = {'a': _A}
_REMPLISSAGES = ('noFill', 'solidFill', 'gradFill', 'blipFill', 'pattFill', 'grpFill')


def _interpoler(t):
    return ''.join(f'{round(a + (b - a) * t):02X}' for a, b in zip(VERT, BLEU))


def _est_blanc(rPr):
    for el in rPr.findall('a:solidFill/a:srgbClr', _NS):
        if el.get('val', '').upper() == 'FFFFFF':
            return True
    for el in rPr.findall('a:solidFill/a:schemeClr', _NS):
        if el.get('val') in ('lt1', 'bg1'):
            return True
    return False


def _colorer(rPr, hexa):
    from lxml import etree
    for tag in _REMPLISSAGES:
        for el in rPr.findall(f'a:{tag}', _NS):
            rPr.remove(el)
    fill = etree.fromstring(
        f'<a:solidFill xmlns:a="{_A}"><a:srgbClr val="{hexa}"/></a:solidFill>')
    ln = rPr.find('a:ln', _NS)
    rPr.insert(list(rPr).index(ln) + 1 if ln is not None else 0, fill)


def _degrader_paragraphe(p):
    from lxml import etree
    runs = p.runs
    if not runs:
        return 0
    texte = ''.join(r.text for r in runs)
    if not texte.strip():
        return 0

    modele = _copy.deepcopy(runs[0]._r)
    rPr_modele = modele.find('a:rPr', _NS)
    if rPr_modele is not None and _est_blanc(rPr_modele):
        return 0

    for r in runs:
        r._r.getparent().remove(r._r)
    fin = p._p.find('a:endParaRPr', _NS)

    n = max(len(texte) - 1, 1)
    for i, car in enumerate(texte):
        run = _copy.deepcopy(modele)
        rPr = run.find('a:rPr', _NS)
        if rPr is None:
            rPr = etree.SubElement(run, f'{{{_A}}}rPr')
            run.insert(0, rPr)
        _colorer(rPr, _interpoler(i / n))
        t = run.find('a:t', _NS)
        if t is None:
            t = etree.SubElement(run, f'{{{_A}}}t')
        t.text = car
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        if fin is not None:
            fin.addprevious(run)
        else:
            p._p.append(run)
    return 1


def appliquer_degrade_titres(prs):
    """Passe les grands titres en degrade vert -> bleu. A appeler EN DERNIER."""
    n = 0
    for slide in prs.slides:
        for sh in walk(slide.shapes):
            titre = sh.name in TITRES_NOMMES or (
                sh.is_placeholder and sh.placeholder_format.idx == 0)
            if not (titre and sh.has_text_frame):
                continue
            for p in sh.text_frame.paragraphs:
                n += _degrader_paragraphe(p)
    return n


import os

ICONES = ('analyse', 'brevet', 'budget', 'delai', 'donnees', 'equipe',
          'feuille_de_route', 'livrable', 'perimetre', 'portefeuille',
          'resultat', 'risque')


def chemin_icone(nom):
    """Chemin du PNG d'une icone Ecosys, relatif au dossier du skill."""
    if nom not in ICONES:
        raise KeyError(f'icone {nom!r} inconnue. Disponibles : {", ".join(ICONES)}')
    return os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                        'assets', 'icones', f'{nom}.png')


def remplacer_icone(slide, nom_forme, nom_icone):
    """Remplace une image du gabarit par une icone Ecosys, au meme emplacement."""
    ancienne = forme(slide, nom_forme)
    gauche, haut = ancienne.left, ancienne.top
    cote = min(ancienne.width, ancienne.height)
    gauche += (ancienne.width - cote) // 2
    haut += (ancienne.height - cote) // 2
    ancienne._element.getparent().remove(ancienne._element)

    nouvelle = slide.shapes.add_picture(chemin_icone(nom_icone), gauche, haut, cote, cote)
    nouvelle.name = nom_forme
    return nouvelle


def supprimer_illustrations_vides(prs):
    """Supprime les emplacements d'image restes vides."""
    n = 0
    for slide in prs.slides:
        for sh in list(slide.shapes):
            if not sh.is_placeholder:
                continue
            try:
                type_ph = str(sh.placeholder_format.type)
            except ValueError:
                continue
            if 'PICTURE' not in type_ph:
                continue
            rempli = sh._element.find(
                './/{http://schemas.openxmlformats.org/drawingml/2006/main}blip') is not None
            if not rempli:
                sh._element.getparent().remove(sh._element)
                n += 1
    return n


BANDE_TITRE_Y = 0.30
BANDE_TITRE_H = 0.70
MARGE_DROITE = 0.60


def largeur_disponible(slide, shape, largeur_slide_po=13.333, garde=0.5):
    """Largeur libre a droite d'une forme, dans sa propre bande verticale."""
    from pptx.util import Emu
    gauche = Emu(shape.left).inches
    haut = Emu(shape.top).inches
    bas = haut + Emu(shape.height).inches
    obstacle = largeur_slide_po - 0.60
    for sh in slide.shapes:
        if sh is shape:
            continue
        try:
            y0 = Emu(sh.top).inches
            y1 = y0 + Emu(sh.height).inches
            x0 = Emu(sh.left).inches
        except TypeError:
            continue
        if y0 < bas and y1 > haut and x0 > gauche + garde:
            obstacle = min(obstacle, x0 - 0.25)
    return max(obstacle - gauche, 3.0)


FACTEUR_GRAS = 0.62
INTERLIGNE = 1.24


def taille_ajustee(texte, largeur_po, hauteur_po, maxi, mini=14):
    """Plus grande taille de police pour laquelle le texte tient dans la boite."""
    taille = maxi
    while taille > mini:
        par_ligne = max(int(largeur_po * 72 / (taille * FACTEUR_GRAS)), 1)
        lignes = max(-(-len(texte) // par_ligne), 1)
        if lignes * taille * INTERLIGNE / 72 <= hauteur_po:
            return taille
        taille -= 2
    return mini


def poser_titre_couverture(slide, shape, titre):
    """Ecrit le titre d'une COUVERTURE ou d'un INTERCALAIRE."""
    from pptx.util import Emu, Inches, Pt
    gauche = Emu(shape.left).inches
    haut = Emu(shape.top).inches
    hauteur = Emu(shape.height).inches
    largeur = largeur_disponible(slide, shape)
    shape.left = Inches(gauche)
    shape.top = Inches(haut)
    shape.height = Inches(hauteur)
    shape.width = Inches(largeur)
    shape.text_frame.word_wrap = True
    taille = taille_ajustee(titre, largeur, hauteur, maxi=40, mini=16)
    ecrire(shape, titre)
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(taille)
    return taille


def poser_message_titre(slide, shape, message, largeur_slide_po=13.333):
    """Ecrit un titre-message et le contraint a la bande reservee par le DA."""
    from pptx.util import Emu, Inches, Pt
    gauche_po = Emu(shape.left).inches
    shape.top = Inches(BANDE_TITRE_Y)
    shape.height = Inches(BANDE_TITRE_H)
    largeur_po = largeur_disponible(slide, shape, largeur_slide_po)
    shape.left = Inches(gauche_po)
    shape.top = Inches(BANDE_TITRE_Y)
    shape.height = Inches(BANDE_TITRE_H)
    shape.width = Inches(largeur_po)
    shape.text_frame.word_wrap = True
    n = len(message)
    taille = 22 if n <= 45 else 18 if n <= 70 else 15
    ecrire(shape, message)
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(taille)
    return taille


ZONE_CONTENU = (0.91, 1.60, 11.50, 4.90)
NAVY = (0x04, 0x00, 0x3E)
VERT = (0x10, 0xBC, 0x4B)
BLEU = (0x00, 0xA3, 0xE9)
GRIS_FOND = (0xF4, 0xF4, 0xF6)
BLANC = (0xFF, 0xFF, 0xFF)


def _rgb(t):
    from pptx.dml.color import RGBColor
    return RGBColor(*t)


def _bloc(slide, forme, x, y, w, h, fond=None, bordure=None, nom=None):
    from pptx.util import Inches
    sh = slide.shapes.add_shape(forme, Inches(x), Inches(y), Inches(w), Inches(h))
    if fond is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = _rgb(fond)
    if bordure is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = _rgb(bordure)
        sh.line.width = Pt_(1.25)
    sh.shadow.inherit = False
    if nom:
        sh.name = nom
    return sh


def Pt_(v):
    from pptx.util import Pt
    return Pt(v)


def _texte(shape, lignes, taille, couleur, gras=False, centre=True, haut=False):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.TOP if haut else MSO_ANCHOR.MIDDLE
    if isinstance(lignes, str):
        lignes = [lignes]
    ecrire(shape, lignes)
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER if centre else PP_ALIGN.LEFT
        for r in p.runs:
            r.font.size = Pt_(taille)
            r.font.bold = gras
            r.font.color.rgb = _rgb(couleur)
            r.font.name = 'Montserrat'


def bandeau_resultats_cles(slide, entrees, zone=None, libelle='Resultats cles'):
    """Bandeau navy a tuiles chiffrees, en bas de slide."""
    from pptx.enum.shapes import MSO_SHAPE
    if not 2 <= len(entrees) <= 4:
        raise ValueError(f'{len(entrees)} tuiles : le bandeau en accepte 2 a 4')
    gx, gy, gw, gh = zone or (ZONE_CONTENU[0], 5.55, ZONE_CONTENU[2], 1.15)
    etq = _bloc(slide, MSO_SHAPE.RECTANGLE, gx, gy - 0.30, 2.2, 0.26)
    _texte(etq, libelle, 11, VERT, gras=True, centre=False)
    fond = _bloc(slide, MSO_SHAPE.ROUNDED_RECTANGLE, gx, gy, gw, gh,
                 fond=NAVY, nom='visuel_resultats_cles')
    fond.adjustments[0] = 0.12
    n = len(entrees)
    marge = 0.20
    largeur = (gw - 2 * marge) / n
    for i, (chiffre, lib) in enumerate(entrees):
        x = gx + marge + i * largeur
        c = _bloc(slide, MSO_SHAPE.RECTANGLE, x, gy + 0.12, largeur, 0.48)
        _texte(c, str(chiffre), 24, VERT, gras=True)
        l = _bloc(slide, MSO_SHAPE.RECTANGLE, x, gy + 0.58, largeur, 0.44)
        _texte(l, lib, 10, BLANC, haut=True)
        if i:
            _bloc(slide, MSO_SHAPE.RECTANGLE, x - 0.005, gy + 0.20, 0.012, 0.72,
                  fond=(0x36, 0x33, 0x65))
    return fond


def pipeline_numerote(slide, etapes, zone=None):
    """Suite de cartes numerotees reliees par des fleches, avec encart de sortie."""
    from pptx.enum.shapes import MSO_SHAPE
    if not 3 <= len(etapes) <= 5:
        raise ValueError(f'{len(etapes)} etapes : le pipeline en accepte 3 a 5')
    gx, gy, gw, gh = zone or ZONE_CONTENU
    n = len(etapes)
    ecart = 0.34
    largeur = (gw - (n - 1) * ecart) / n
    plus_long = max((len(e.get('detail') or '') for e in etapes), default=0)
    par_ligne = max(int(largeur * 72 / (9 * 0.52)), 1)
    lignes = max(-(-plus_long // par_ligne), 1)
    h_carte = min(max(0.95 + lignes * 0.16, 1.4), gh - 1.05)
    h_totale = 0.22 + h_carte + 0.12 + 0.62
    gy = gy + max((gh - h_totale) / 2, 0)
    for i, e in enumerate(etapes):
        x = gx + i * (largeur + ecart)
        carte = _bloc(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, gy + 0.22, largeur, h_carte,
                      fond=GRIS_FOND, nom=f'visuel_etape{i + 1}')
        carte.adjustments[0] = 0.08
        titre = _bloc(slide, MSO_SHAPE.RECTANGLE, x + 0.08, gy + 0.46, largeur - 0.16, 0.42)
        _texte(titre, e['titre'], 11, NAVY, gras=True)
        if e.get('detail'):
            det = _bloc(slide, MSO_SHAPE.RECTANGLE, x + 0.08, gy + 0.90,
                        largeur - 0.16, h_carte - 0.80)
            _texte(det, e['detail'], 9, (0x33, 0x33, 0x33), centre=False, haut=True)
        d = 0.42
        past = _bloc(slide, MSO_SHAPE.OVAL, x + largeur / 2 - d / 2, gy, d, d, fond=NAVY)
        _texte(past, str(i + 1), 13, BLANC, gras=True)
        if e.get('sortie'):
            out = _bloc(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, gy + 0.22 + h_carte + 0.12,
                        largeur, 0.62, fond=VERT)
            out.adjustments[0] = 0.14
            _texte(out, e['sortie'], 10, BLANC, gras=True)
        if i < n - 1:
            fl = _bloc(slide, MSO_SHAPE.RIGHT_ARROW, x + largeur + 0.06,
                       gy + 0.22 + h_carte / 2 - 0.11, ecart - 0.12, 0.22, fond=BLEU)
            fl.name = f'visuel_fleche{i + 1}'
    return n


def rotation_variantes(slides_par_theme, nombre, theme_prefere=None):
    """Choisit nombre slides sources en faisant TOURNER les variantes."""
    items = list(slides_par_theme.items())
    ordonne = []
    if theme_prefere:
        for cle, val in items:
            if theme_prefere.lower() in cle.lower():
                ordonne.append((cle, val))
    ordonne += [kv for kv in items if kv not in ordonne]
    plat = []
    for _, val in ordonne:
        plat += val if isinstance(val, list) else [val]
    if not plat:
        raise ValueError('aucune variante disponible pour ce gabarit')
    if nombre > len(plat):
        print(f'  WARN {nombre} slides demandees pour {len(plat)} variantes : '
              'certaines photos se repeteront')
    return [plat[i % len(plat)] for i in range(nombre)]


PROGRESSION = ((0x10, 0xBC, 0x4B), (0x08, 0xB0, 0x9A), (0x00, 0xA3, 0xE9), (0x36, 0x33, 0x65))


def couches_empilees(slide, couches, zone=None):
    """Bandes horizontales empilees : un verbe colore a gauche, le detail a droite."""
    from pptx.enum.shapes import MSO_SHAPE
    if not 2 <= len(couches) <= 4:
        raise ValueError(f'{len(couches)} couches : le motif en accepte 2 a 4')
    gx, gy, gw, gh = zone or ZONE_CONTENU
    n = len(couches)
    ecart = 0.16
    h = min((gh - (n - 1) * ecart) / n, 1.30)
    gy = gy + max((gh - (n * h + (n - 1) * ecart)) / 2, 0)
    l_verbe = 2.70
    for i, c in enumerate(couches):
        y = gy + i * (h + ecart)
        teinte = PROGRESSION[i % len(PROGRESSION)]
        fond = _bloc(slide, MSO_SHAPE.ROUNDED_RECTANGLE, gx, y, gw, h,
                     fond=GRIS_FOND, nom=f'visuel_couche{i + 1}')
        fond.adjustments[0] = 0.08
        etq = _bloc(slide, MSO_SHAPE.ROUNDED_RECTANGLE, gx, y, l_verbe, h, fond=teinte)
        etq.adjustments[0] = 0.08
        _texte(etq, c['verbe'], 14, BLANC, gras=True)
        det = _bloc(slide, MSO_SHAPE.RECTANGLE, gx + l_verbe + 0.22, y + 0.06,
                    gw - l_verbe - 0.44, h - 0.12)
        _texte(det, c['detail'], 11, (0x33, 0x33, 0x33), centre=False)
    return n


def trois_icones(slide, volets, zone=None, accroche=None):
    """Icones Ecosys en ligne, titre et texte dessous, bandeau de synthese optionnel."""
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    if not 2 <= len(volets) <= 4:
        raise ValueError(f'{len(volets)} volets : le motif en accepte 2 a 4')
    gx, gy, gw, gh = zone or ZONE_CONTENU
    n = len(volets)
    ecart = 0.45
    largeur = (gw - (n - 1) * ecart) / n
    d = 1.05
    h_bandeau = 0.62 if accroche else 0.0
    plus_long = max((len(v.get('texte') or '') for v in volets), default=0)
    par_ligne = max(int((largeur - 0.20) * 72 / (10 * 0.52)), 12)
    n_lignes = max(-(-plus_long // par_ligne), 1)
    h_texte = n_lignes * 0.17 + 0.10
    h_bloc = d + 0.62 + h_texte
    h_totale = h_bloc + (0.18 + h_bandeau if accroche else 0)
    gy = gy + max((gh - h_totale) / 2, 0.10)
    for i, v in enumerate(volets):
        x = gx + i * (largeur + ecart)
        cx = x + largeur / 2
        past = _bloc(slide, MSO_SHAPE.OVAL, cx - d / 2, gy, d, d, fond=NAVY,
                     nom=f'visuel_pastille{i + 1}')
        marge = d * 0.24
        slide.shapes.add_picture(chemin_icone(v['icone']),
                                 Inches(cx - d / 2 + marge), Inches(gy + marge),
                                 Inches(d - 2 * marge), Inches(d - 2 * marge))
        tit = _bloc(slide, MSO_SHAPE.RECTANGLE, x, gy + d + 0.16, largeur, 0.46)
        _texte(tit, v['titre'], 13, NAVY, gras=True)
        if v.get('texte'):
            txt = _bloc(slide, MSO_SHAPE.RECTANGLE, x + 0.10, gy + d + 0.62,
                        largeur - 0.20, h_texte)
            _texte(txt, v['texte'], 10, (0x33, 0x33, 0x33), haut=True)
        _bloc(slide, MSO_SHAPE.RECTANGLE, cx - 0.22, gy + d + 0.10, 0.44, 0.045,
              fond=PROGRESSION[i % len(PROGRESSION)])
    if accroche:
        band = _bloc(slide, MSO_SHAPE.ROUNDED_RECTANGLE, gx, gy + h_bloc + 0.18,
                     gw, h_bandeau, fond=VERT, nom='visuel_accroche')
        band.adjustments[0] = 0.14
        _texte(band, accroche, 13, BLANC, gras=True)
    return n
